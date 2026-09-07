from pathlib import Path
import json,re,html,shutil
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'build/generated';OUT.mkdir(parents=True,exist_ok=True)
PKG=ROOT/'research';GENERATED=OUT/'research';GENERATED.mkdir(exist_ok=True)
source=(ROOT/'authoring/manuscript_body.md').read_text()
metadata=json.loads((ROOT/'authoring/reference_metadata.json').read_text())
dois=['10.1136/bmj-2025-086295','10.1038/s44220-026-00656-y','10.1038/s41586-026-10910-z','10.1038/s41467-026-71962-3','10.1093/brain/awag206','10.3389/fpsyt.2025.1505297','10.1002/aur.2954','10.1017/S0033291724001387','10.1038/s41591-024-03196-z','10.1186/s12888-025-06689-4','10.1016/j.neubiorev.2024.105807','10.1016/S0006-3495(72)86068-5','10.1103/PhysRevLett.88.174102','10.1007/s12559-021-09931-9','10.1126/science.abd0380','10.1177/1362361318786721','10.1371/journal.pcbi.1003285','10.3389/fnhum.2014.00020','10.1038/mp.2013.78','10.1038/sdata.2017.10','10.3389/fnins.2018.00477','10.1038/s41598-023-32649-7']
year_overrides={'10.1007/s12559-021-09931-9':2023,'10.1177/1362361318786721':2019,'10.1038/mp.2013.78':2014}
journals={'Frontiers in Neuroscience':'Front. Neurosci.','Frontiers in Human Neuroscience':'Front. Hum. Neurosci.','Frontiers in Psychiatry':'Front. Psychiatry','Nature Mental Health':'Nat. Ment. Health','Nature Communications':'Nat. Commun.','Nature Medicine':'Nat. Med.','Scientific Reports':'Sci. Rep.','Autism Research':'Autism Res.','Psychological Medicine':'Psychol. Med.','Neuroscience & Biobehavioral Reviews':'Neurosci. Biobehav. Rev.','Biophysical Journal':'Biophys. J.','Physical Review Letters':'Phys. Rev. Lett.','Cognitive Computation':'Cogn. Comput.','PLoS Computational Biology':'PLoS Comput. Biol.','Molecular Psychiatry':'Mol. Psychiatry','Scientific Data':'Sci. Data'}
def plain(s):return html.unescape(re.sub('<[^>]+>','',s))
def initials(g):return ' '.join(x[0]+'.' for x in re.findall(r'[^\W\d_]+',g,flags=re.UNICODE))
refs=[];bib=[]
for doi in dois:
    m=metadata[doi];authors=m['author'];year=year_overrides.get(doi,m['published']['date-parts'][0][0]);an=[a['family']+', '+initials(a.get('given','')) for a in authors[:6]]
    astr=', '.join(an)+(', et al.' if len(authors)>6 else '')
    title=plain(m['title'][0]).rstrip('.');journal=plain(m['container-title'][0]);volume=m.get('volume','');loc=m.get('page') or m.get('article-number','')
    if doi=='10.3389/fnhum.2014.00020':loc='20'
    if doi=='10.1371/journal.pcbi.1003285':loc='e1003285'
    if doi=='10.1093/brain/awag206':loc='awag206 (advance online publication)'
    suffix=(' '+volume if volume else '')+(', '+loc if loc else '')
    ref=f'{astr} ({year}). {title}. {journals.get(journal,journal)}{suffix}. doi: {doi}'
    refs.append((authors[0]['family'],year,ref))
    key=re.sub(r'[^A-Za-z]','',authors[0]['family'])+str(year)
    bib.append('@article{'+key+',\n  author = {'+' and '.join(a['family']+', '+a.get('given','') for a in authors)+'},\n  title = {'+title+'},\n  journal = {'+journal+'},\n  year = {'+str(year)+'},\n  volume = {'+volume+'},\n  pages = {'+loc+'},\n  doi = {'+doi+'}\n}')
refs.sort(key=lambda x:(x[0].lower(),x[1]))
(GENERATED/'references.bib').write_text('\n\n'.join(bib))
(GENERATED/'references.txt').write_text('\n\n'.join(r[2] for r in refs))
(GENERATED/'manuscript_source.md').write_text(source+'\n\n## References\n\n'+'\n\n'.join(r[2] for r in refs))

doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.8);sec.left_margin=Inches(1);sec.right_margin=Inches(.85)
sec.header_distance=sec.footer_distance=Inches(.35)
styles=doc.styles
for st in styles:
    for node in st.element.xpath('.//w:pBdr'):node.getparent().remove(node)
    for node in st.element.xpath('.//w:rFonts'):
        for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
            node.attrib.pop(qn('w:'+attr),None)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    styles[name].font.name='Times New Roman';styles[name].font.color.rgb=RGBColor(0,0,0)
normal=styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.0;normal.paragraph_format.widow_control=True
for name,size,before,after in [('Title',18,0,12),('Heading 1',13,14,6),('Heading 2',11.5,10,5),('Heading 3',11,8,4)]:
    s=styles[name];s.font.size=Pt(size);s.font.bold=True;s.paragraph_format.space_before=Pt(before);s.paragraph_format.space_after=Pt(after);s.paragraph_format.keep_with_next=True
styles['Caption'].font.size=Pt(10);styles['Caption'].font.italic=False;styles['Caption'].font.bold=False
for name in ['Reference','Table Text','Equation']:
    if name not in styles:styles.add_style(name,1)
    styles[name].base_style=normal;styles[name].font.name='Times New Roman';styles[name].font.size=Pt(10)
styles['Reference'].paragraph_format.left_indent=Inches(.2);styles['Reference'].paragraph_format.first_line_indent=Inches(-.2)
styles['Table Text'].paragraph_format.space_after=Pt(3)
styles['Equation'].font.name='Cambria Math';styles['Equation'].font.size=Pt(10.5);styles['Equation'].paragraph_format.space_after=Pt(4)
for st in ('Header','Footer'):
    sup=OxmlElement('w:suppressLineNumbers');styles[st].element.get_or_add_pPr().append(sup)
line=OxmlElement('w:lnNumType');line.set(qn('w:countBy'),'1');line.set(qn('w:start'),'1');line.set(qn('w:restart'),'continuous');line.set(qn('w:distance'),'240');sec._sectPr.append(line)
h=sec.header.paragraphs[0];sup=OxmlElement('w:suppressLineNumbers');h._p.get_or_add_pPr().append(sup);h.text='Hopkins | Entropy-informed feedback benchmark';h.style='Normal';h.runs[0].font.size=Pt(9);h.runs[0].font.color.rgb=RGBColor(90,90,90)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
f.add_run('Research manuscript draft • 7 September 2026 | ').font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
title=source.splitlines()[0][2:];doc.add_paragraph(title,'Title')
doc.add_paragraph('Ryan Hopkins').runs[0].bold=True
doc.add_paragraph('Affiliation and correspondence: to be confirmed by the author before submission.').runs[0].italic=True
main=source[source.index('## 1 Introduction'):source.index('## Data and code availability')]
wordcount=len(re.findall(r"\b[\w’−-]+\b",main))
abstract=source[source.index('## Abstract')+len('## Abstract'):source.index('## 1 Introduction')].strip()
p=doc.add_paragraph(f'Proposed journal: Frontiers in Computational Neuroscience | Article type: Original Research\nMain-text word count: {wordcount:,} (excluding abstract, references, captions and tables)\nAbstract: {len(abstract.split())} words | Figures: 4 | Tables: 2 | References: {len(refs)}')
for r in p.runs:r.font.size=Pt(9.5)
doc.add_paragraph('Keywords: autism; closed-loop neuromodulation; brain–computer interface; permutation entropy; neural-mass modeling; computational reproducibility; person-specific feedback').runs[0].font.size=Pt(10)
for block in source[source.index('## Abstract'):].split('\n\n'):
    block=block.strip()
    if not block:continue
    if block.startswith('### '):doc.add_paragraph(block[4:],'Heading 2')
    elif block.startswith('## '):doc.add_paragraph(block[3:],'Heading 1')
    elif block.startswith('@EQUATION '):
        p=doc.add_paragraph(style='Equation');eq=OxmlElement('m:oMath')
        def mr(text):
            r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;t.set(qn('xml:space'),'preserve');r.append(t);return r
        text=block[len('@EQUATION '):].replace('W^0_ij','W_ij⁰')
        pattern=re.compile(r'([A-Za-zτξΣĒ]+)_([A-Za-z]+(?:,[ij])?)')
        pos=0
        for match in pattern.finditer(text):
            if match.start()>pos:eq.append(mr(text[pos:match.start()]))
            sub=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(mr(match[1]));su=OxmlElement('m:sub');su.append(mr(match[2]));sub.extend([e,su]);eq.append(sub);pos=match.end()
        if pos<len(text):eq.append(mr(text[pos:]))
        p._p.append(eq)
    else:
        p=doc.add_paragraph();last=0
        for match in re.finditer(r'\b(E|I|W|M|H|p|mean|Σ)_([ijk]+)\b',block):
            p.add_run(block[last:match.start()]+match[1]);r=p.add_run(match[2]);r.font.subscript=True;last=match.end()
        p.add_run(block[last:])
doc.add_paragraph('References','Heading 1')
for _,_,ref in refs:doc.add_paragraph(ref,'Reference')

tables=[('Table 1. Model and analysis parameters. All inputs and activities are dimensionless unless a time unit is given. Ranges define synthetic variation, not autistic population estimates.',
['Component','Value / implementation'],[
['Ensemble','24 individuals; 16 E/I populations each; 3 evaluation seeds'],['Integration / observation','Euler 2 ms; output 100 Hz; 60 s per run'],['Time constants','Excitation 40 ms; inhibition 80 ms; current noise 150 ms'],['Gain / coupling / drive','Uniform 1.25–1.75 / 1.4–2.6 / 0.55–1.05'],['Local recurrent weights','wEE uniform 2.5–4.5; wEI uniform 9–11; I coefficients 10 and 2'],['Disturbance','Random sign × uniform 0.35–0.80; second context multiplies by −0.6'],['Intrinsic current noise','Ornstein–Uhlenbeck scale 0.04'],['Sensor noise / artifact','Nominal SD 0.001; stress SD 0.05; artifact 0.20 × u × sin(2π·37t)'],['Permutation entropy','Dimension 3; lag 20 ms; 1-s windows; normalization ln(6)'],['Calibration','Separate reference; local entropy slope from ±0.20 inputs'],['Entropy abstention / correction','|slope| < 0.01; otherwise correction clipped to ±0.35'],['Feedback gains','Rate: 4; combined: rate + 0.3 × entropy correction'],['Actuator constraints','Absolute cap 0.60; slew 0.15 per 1-s update'],['Fixed comparator','7 candidates, −0.60 to +0.60, selected using an independent perturbed seed'],['Readout','12 non-input populations; ridge penalty 10; target lag 150 ms'],['Plasticity (optional)','η = 0.15; activity mean 2 s; relaxation 20 s; bounds [0.25W⁰, 2W⁰]'],['Uncertainty','4,000 bootstrap draws; seeds averaged within individuals']]),
('Table 2. Nominal engineering outcomes. Values are means with 95% percentile bootstrap intervals over 24 synthetic individuals after averaging three seeds. These are not clinical outcome estimates.',
['Strategy','Tracking MSE','Accuracy (%)','Input energy'],[])]
study=json.loads((PKG/'results/summary.json').read_text())
for arm,label in [('sham','No input'),('fixed','Fixed input'),('rate','Rate feedback'),('entropy','Entropy feedback'),('combined','Rate + entropy'),('yoked','Yoked replay')]:
    v=study['summary']['nominal_'+arm];row=[label]
    for key,fac,d in [('tracking_mse',1,2),('accuracy',100,2),('control_energy',1,4)]:
        m,lo,hi=v[key];row.append(f'{m*fac:.{d}f}\n[{lo*fac:.{d}f}, {hi*fac:.{d}f}]')
    tables[1][2].append(row)
for caption,head,rows in tables:
    doc.add_page_break();doc.add_paragraph(caption,'Caption');doc.add_paragraph('').paragraph_format.space_after=Pt(0)
    table=doc.add_table(rows=1,cols=len(head));table.autofit=False
    widths=[2,4.65] if len(head)==2 else [1.5,1.75,1.7,1.7]
    for column,width in zip(table.columns,widths):column.width=Inches(width)
    for cell,width in zip(table.rows[0].cells,widths):cell.width=Inches(width)
    for i,x in enumerate(head):table.cell(0,i).text=x
    for row in rows:
        cells=table.add_row().cells
        for i,x in enumerate(row):cells[i].text=x;cells[i].width=Inches(widths[i])
    for ridx,row in enumerate(table.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            for p in cell.paragraphs:
                p.style='Table Text'
                if ridx==0:
                    for r in p.runs:r.bold=True
            if ridx==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E9EEF2');cell._tc.get_or_add_tcPr().append(sh)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    doc.add_paragraph('')

figs=[('Figure_1_Control_Architecture','Figure 1. Measurement, feedback and independent evaluation are separated. The synthetic circuit generates an observation; feedback uses reference errors; numerical constraints bound the returned input. A separately trained task readout evaluates performance. Anatomy and therapeutic efficacy are not inferred from this architecture.','Control diagram with a synthetic neural circuit, noisy observation, feedback rule and bounded actuator in a loop; independent task readout is the outcome.'),
('Figure_2_Benchmark','Figure 2. Nominal controller outcomes. (A) Mean task-tracking squared error. (B) Binary task accuracy, with 50% shown as a chance reference and the held-out unperturbed accuracy interval shaded. Point estimates and intervals are based on 24 synthetic individuals, each averaged over three seeds. Combined feedback lowers tracking error relative to rate feedback without recovering useful sign discrimination. Fixed-input calibration access differs from that of adaptive rules.','Six controller strategies compared by tracking error and accuracy. Combined error is 76.83 versus 81.36 for rate feedback; both accuracies are approximately 50%, below the 80.27% unperturbed reference.'),
('Figure_3_Robustness','Figure 3. Incremental tracking error from adding entropy feedback. Negative values favor combined feedback over rate feedback; positive values favor rate feedback. Clean observations and optional plasticity have negative intervals, sensor noise has a positive interval, and delay and artifact intervals span zero. The four stress conditions were tested separately. Intervals describe synthetic ensemble variation.','Forest plot: combined minus rate error is −4.53 clean, −0.77 delayed, +8.28 noisy, +3.52 with artifact, and −4.60 with plasticity. Noise reverses the clean-measurement result.'),
('Figure_4_Entropy_Diagnostics','Figure 4. Exploratory entropy diagnostics. (A) Each line connects 11 gain settings for one of eight synthetic individuals. No optimal or inverted-U entropy curve was fitted. (B) Original whole-record permutation entropy is compared with the mean and range of 40 phase surrogates preserving Fourier magnitudes. Whole-record estimates differ from the mean short-window estimator used online. These descriptive diagnostics do not establish biomarker specificity.','Two panels show varied entropy–accuracy trajectories as gain changes and original versus phase-randomized entropy across eight synthetic individuals.')]
for stem,caption,alt in figs:
    doc.add_page_break();shape=doc.add_picture(str(PKG/'figures'/f'{stem}.png'),width=Inches(6.65));shape._inline.docPr.set('descr',alt);doc.add_paragraph(caption,'Caption')
doc.core_properties.title=title;doc.core_properties.author='Ryan Hopkins';doc.core_properties.subject='Exploratory computational neuroscience manuscript; author verification required before submission';doc.core_properties.keywords='autism, entropy, BCI, neural mass, feedback'
doc.save(OUT/'Hopkins_Entropy_Feedback_Manuscript.docx')
(OUT/'manuscript_counts.json').write_text(json.dumps({'main_words':wordcount,'abstract_words':len(abstract.split()),'references':len(refs)}))
print('Saved manuscript.',wordcount,'main-text words;',len(refs),'references.')
